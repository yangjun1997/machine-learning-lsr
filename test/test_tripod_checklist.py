from docx import Document

from mechine_learning_lsr.tripod_checklist import (
    CALCULATOR_URL,
    CHECKLIST_ITEMS,
    REPOSITORY_URL,
    generate_checklist,
)


def test_tripod_ai_checklist_is_complete_and_explicit(tmp_path):
    output = generate_checklist(tmp_path / "TRIPOD_AI_Checklist.docx")
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert len(CHECKLIST_ITEMS) == 52
    assert "TRIPOD+AI Checklist" in text
    assert "10.1136/bmj-2023-078378" in text
    assert len(document.tables[0].rows) == 53
    assert "Provide details of the availability of the analytical code" in table_text
    assert REPOSITORY_URL in table_text
    assert CALCULATOR_URL in table_text
    assert "Not reported" in table_text
    assert "author input" in table_text.lower()
