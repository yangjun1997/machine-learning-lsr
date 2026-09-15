from pathlib import Path

from docx import Document

from mechine_learning_lsr.word_export import (
    FIGURES,
    SUPPLEMENTARY_FIGURE_CAPTIONS,
    TABLES,
    build_supplementary_figure_legends_document,
    build_word_document,
)


def _markdown_row_count(path: Path) -> int:
    return sum(
        line.startswith("|") and "---" not in line
        for line in path.read_text(encoding="utf-8").splitlines()
    )


def test_word_export_embeds_main_figures_and_editable_tables(tmp_path):
    output = build_word_document(output=tmp_path / "manuscript.docx")
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    assert len(document.inline_shapes) == len(FIGURES) == 5
    assert len(document.tables) == len(TABLES) == 3
    assert len(document.tables[0].columns) == 5
    assert [len(table.rows) for table in document.tables] == [
        _markdown_row_count(TABLES[number]) for number in TABLES
    ]
    assert "Figure 5. Exploratory residual-LSR effect modification" in text
    assert "AUROC" in document.tables[2].cell(1, 0).text
    assert document.tables[0].cell(0, 4).text == "Statistical method"
    assert "[[AUTHOR INPUT REQUIRED:" in text
    assert "Open Table" not in text
    assert "M1" not in text
    assert "lsr_combo" not in text
    assert "Logistic regression (L1)" not in table_text
    assert "Logistic regression (L2)" not in table_text
    assert "Logistic regression (elastic net)" in table_text


def test_supplementary_figure_legends_export(tmp_path):
    output = build_supplementary_figure_legends_document(tmp_path / "supplementary_legends.docx")
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert len(SUPPLEMENTARY_FIGURE_CAPTIONS) == 5
    assert "Figure S1. Group 1 learning curve" in text
    assert "Figure S5. Confusion matrix" in text
    assert "300 refits" in text
    assert "2,000 patient-level bootstrap" in text
