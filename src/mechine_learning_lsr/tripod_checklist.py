from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
import os
import shutil

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE_URL = "https://www.tripod-statement.org/wp-content/uploads/2019/12/TRIPODAI_checklist.pdf"
ARTICLE_URL = "https://doi.org/10.1136/bmj-2023-078378"
REPOSITORY_URL = "https://github.com/yangjun1997/machine-learning-lsr"
CALCULATOR_URL = "https://hfs-mvd-risk-calculator.streamlit.app/"
MANUSCRIPT_PATH = Path("results/Manuscripts.docx")
CHECKLIST_PATH = Path("results/TRIPOD_AI_Checklist.docx")
BACKUP_PATH = Path("results/Manuscripts_pre_TRIPOD_AI.docx")
LEGACY_CALCULATOR_NOTE = (
    "[[AUTHOR INPUT REQUIRED: insert the confirmed public Streamlit calculator URL; "
    "the previously supplied URL returned HTTP 404 on 20 September 2026.]]"
)
CALCULATOR_NOTE = f"the web-based research calculator is publicly accessible at {CALCULATOR_URL}"


@dataclass(frozen=True)
class ChecklistItem:
    section: str
    topic: str
    item: str
    applicability: str
    text: str


_ITEM_ROWS = """
Title|Title|1|D;E|Identify the study as developing or evaluating the performance of a multivariable prediction model, the target population, and the outcome to be predicted
Abstract|Abstract|2|D;E|See TRIPOD+AI for Abstracts checklist
Introduction|Background|3a|D;E|Explain the healthcare context (including whether diagnostic or prognostic) and rationale for developing or evaluating the prediction model, including references to existing models
Introduction|Background|3b|D;E|Describe the target population and the intended purpose of the prediction model in the context of the care pathway, including its intended users (eg, healthcare professionals, patients, public)
Introduction|Background|3c|D;E|Describe any known health inequalities between sociodemographic groups
Introduction|Objectives|4|D;E|Specify the study objectives, including whether the study describes the development or validation of a prediction model (or both)
Methods|Data|5a|D;E|Describe the sources of data separately for the development and evaluation datasets (eg, randomised trial, cohort, routine care or registry data), the rationale for using these data, and representativeness of the data
Methods|Data|5b|D;E|Specify the dates of the collected participant data, including start and end of participant accrual; and, if applicable, end of follow-up
Methods|Participants|6a|D;E|Specify key elements of the study setting (eg, primary care, secondary care, general population) including the number and location of centres
Methods|Participants|6b|D;E|Describe the eligibility criteria for study participants
Methods|Participants|6c|D;E|Give details of any treatments received, and how they were handled during model development or evaluation, if relevant
Methods|Data preparation|7|D;E|Describe any data pre-processing and quality checking, including whether this was similar across relevant sociodemographic groups
Methods|Outcome|8a|D;E|Clearly define the outcome that is being predicted and the time horizon, including how and when assessed, the rationale for choosing this outcome, and whether the method of outcome assessment is consistent across sociodemographic groups
Methods|Outcome|8b|D;E|If outcome assessment requires subjective interpretation, describe the qualifications and demographic characteristics of the outcome assessors
Methods|Outcome|8c|D;E|Report any actions to blind assessment of the outcome to be predicted
Methods|Predictors|9a|D|Describe the choice of initial predictors (eg, literature, previous models, all available predictors) and any pre-selection of predictors before model building
Methods|Predictors|9b|D;E|Clearly define all predictors, including how and when they were measured (and any actions to blind assessment of predictors for the outcome and other predictors)
Methods|Predictors|9c|D;E|If predictor measurement requires subjective interpretation, describe the qualifications and demographic characteristics of the predictor assessors
Methods|Sample size|10|D;E|Explain how the study size was arrived at (separately for development and evaluation), and justify that the study size was sufficient to answer the research question. Include details of any sample size calculation
Methods|Missing data|11|D;E|Describe how missing data were handled. Provide reasons for omitting any data
Methods|Analytical methods|12a|D|Describe how the data were used (eg, for development and evaluation of model performance) in the analysis, including whether the data were partitioned, considering any sample size requirements
Methods|Analytical methods|12b|D|Depending on the type of model, describe how predictors were handled in the analyses (functional form, rescaling, transformation, or any standardisation)
Methods|Analytical methods|12c|D|Specify the type of model, rationale, all model building steps, including any hyperparameter tuning, and method for internal validation
Methods|Analytical methods|12d|D;E|Describe if and how any heterogeneity in estimates of model parameter values and model performance was handled and quantified across clusters (eg, hospitals, countries). See TRIPOD-Cluster for additional considerations
Methods|Analytical methods|12e|D;E|Specify all measures and plots used (and their rationale) to evaluate model performance (eg, discrimination, calibration, clinical utility) and, if relevant, to compare multiple models
Methods|Analytical methods|12f|E|Describe any model updating (eg, recalibration) arising from the model evaluation, either overall or for particular sociodemographic groups or settings
Methods|Analytical methods|12g|E|For model evaluation, describe how the model predictions were calculated (eg, formula, code, object, application programming interface)
Methods|Class imbalance|13|D;E|If class imbalance methods were used, state why and how this was done, and any subsequent methods to recalibrate the model or the model predictions
Methods|Fairness|14|D;E|Describe any approaches that were used to address model fairness and their rationale
Methods|Model output|15|D|Specify the output of the prediction model (eg, probabilities, classification). Provide details and rationale for any classification and how the thresholds were identified
Methods|Training versus evaluation|16|D;E|Identify any differences between the development and evaluation data in healthcare setting, eligibility criteria, outcome, and predictors
Methods|Ethical approval|17|D;E|Name the institutional research board or ethics committee that approved the study and describe the participant informed consent or the ethics committee waiver of informed consent
Open science|Funding|18a|D;E|Give the source of funding and the role of the funders for the present study
Open science|Conflicts of interest|18b|D;E|Declare any conflicts of interest and financial disclosures for all authors
Open science|Protocol|18c|D;E|Indicate where the study protocol can be accessed or state that a protocol was not prepared
Open science|Registration|18d|D;E|Provide registration information for the study, including register name and registration number, or state that the study was not registered
Open science|Data sharing|18e|D;E|Provide details of the availability of the study data
Open science|Code sharing|18f|D;E|Provide details of the availability of the analytical code
Patient and public involvement|Patient and public involvement|19|D;E|Provide details of any patient and public involvement during the design, conduct, reporting, interpretation, or dissemination of the study or state no involvement
Results|Participants|20a|D;E|Describe the flow of participants through the study, including the number of participants with and without the outcome and, if applicable, a summary of the follow-up time. A diagram may be helpful
Results|Participants|20b|D;E|Report the characteristics overall and, where applicable, for each data source or setting, including the key dates, key predictors (including demographics), treatments received, sample size, number of outcome events, follow-up time, and amount of missing data. A table may be helpful. Report any differences across key demographic groups
Results|Participants|20c|E|For model evaluation, show a comparison with the development data of the distribution of important predictors (demographics, predictors, and outcome)
Results|Model development|21|D;E|Specify the number of participants and outcome events in each analysis (eg, for model development, hyperparameter tuning, model evaluation)
Results|Model specification|22|D|Provide details of the full prediction model (eg, formula, code, object, application programming interface) to allow predictions in new individuals and to enable third party evaluation and implementation, including any restrictions to access or reuse (eg, freely available, proprietary)
Results|Model performance|23a|D;E|Report model performance estimates with confidence intervals, including for any key subgroups (eg, sociodemographic). Consider plots to aid presentation
Results|Model performance|23b|D;E|If examined, report results of any heterogeneity in model performance across clusters. See TRIPOD-Cluster for additional details
Results|Model updating|24|E|Report the results from any model updating, including the updated model and subsequent performance
Discussion|Interpretation|25|D;E|Give an overall interpretation of the main results, including issues of fairness in the context of the objectives and previous studies
Discussion|Limitations|26|D;E|Discuss any limitations of the study (such as a non-representative sample, sample size, overfitting, missing data) and their effects on any biases, statistical uncertainty, and generalisability
Discussion|Usability of the model in the context of current care|27a|D|Describe how poor quality or unavailable input data (eg, predictor values) should be assessed and handled when implementing the prediction model
Discussion|Usability of the model in the context of current care|27b|D|Specify whether users will be required to interact in the handling of the input data or use of the model, and what level of expertise is required of users
Discussion|Usability of the model in the context of current care|27c|D;E|Discuss any next steps for future research, with a specific view to applicability and generalisability of the model
""".strip()

CHECKLIST_ITEMS = tuple(ChecklistItem(*row.split("|", 4)) for row in _ITEM_ROWS.splitlines())
assert len(CHECKLIST_ITEMS) == 52


LOCATIONS = {
    "1": "Not reported: the current manuscript file does not contain a title.",
    "2": "Not reported: the current manuscript file does not contain an abstract.",
    "3a": "Introduction, paragraphs 1-4.",
    "3b": "Introduction, final paragraph; Discussion, clinical-implications paragraph. Intended users are clinicians; public/patient use is not proposed.",
    "3c": "Not reported: known health inequalities are not described.",
    "4": "Introduction, final paragraph.",
    "5a": "Methods—Study Design and Participants; Methods—Temporal Split and Data Quality.",
    "5b": "Methods—Temporal Split and Data Quality. End of follow-up is reported as a minimum of one year, but the final follow-up date is not reported.",
    "6a": "Methods—Study Design and Participants (single secondary-care neurosurgical centre, Beijing Tiantan Hospital).",
    "6b": "Methods—Study Design and Participants.",
    "6c": "Methods—Study Design and Participants; Methods—Outcome and Predictors. All participants underwent MVD; prior botulinum toxin and acupuncture were candidate predictors.",
    "7": "Partly reported: Methods—Outcome and Predictors; Model Development and Freezing. Preprocessing is described, but quality checking across sociodemographic groups is not.",
    "8a": "Methods—Outcome and Predictors.",
    "8b": "Not reported: outcome-assessor qualifications and demographic characteristics are not described.",
    "8c": "Not reported: outcome-assessment blinding is not described.",
    "9a": "Partly reported: Methods—Outcome and Predictors lists the 12 initial predictors; the rationale for choosing all initial predictors is not fully described.",
    "9b": "Partly reported: Methods—Outcome and Predictors defines the predictors and LSR categories; timing, assessor blinding, and detailed measurement procedures are incomplete.",
    "9c": "Not reported: qualifications and demographic characteristics of predictor assessors are not described.",
    "10": "Partly reported: Discussion—Limitations cites the validation-event limitation and an AUROC sample-size reference; a development-sample calculation and complete evaluation-sample justification are not reported.",
    "11": "Methods—Study Design and Participants; Model Development and Freezing. Incomplete records were excluded and numerical imputation is described; missingness amounts and omission reasons are not fully reported.",
    "12a": "Methods—Temporal Split and Data Quality; Model Development and Freezing; Internal Temporal Validation and Statistical Analysis.",
    "12b": "Methods—Outcome and Predictors; Model Development and Freezing.",
    "12c": "Methods—Model Development and Freezing.",
    "12d": "Not applicable: this is a single-centre study and no cluster-level heterogeneity analysis was performed.",
    "12e": "Methods—Model Development and Freezing; Internal Temporal Validation and Statistical Analysis.",
    "12f": "Methods—Internal Temporal Validation and Statistical Analysis (no Group 2 recalibration was performed).",
    "12g": "Methods—Internal Temporal Validation and Statistical Analysis; Code Availability; frozen pipeline object in the repository.",
    "13": "Not applicable to the frozen index model: no class-imbalance method or post-hoc recalibration was used.",
    "14": "Not reported: no model-fairness method or rationale is described.",
    "15": "Methods—Internal Temporal Validation and Statistical Analysis; Results—internal temporal validation. Output is probability; threshold analyses are explicitly exploratory.",
    "16": "Results—Cohort Flow and Characteristics; Table 1.",
    "17": "Methods—Study Design and Participants. Ethics committee name and consent are reported; approval number is not reported.",
    "18a": "Not reported: funding source and funder role are absent; author input is required.",
    "18b": "Not reported: conflicts of interest and financial disclosures are absent; author input is required.",
    "18c": "Not reported: protocol availability or absence is not stated; author input is required.",
    "18d": "Not reported: study registration or non-registration is not stated; author input is required.",
    "18e": "Data Availability statement: participant-level data are not publicly available because of patient privacy and institutional ethics requirements.",
    "18f": f"Code Availability statement; repository: {REPOSITORY_URL}; calculator: {CALCULATOR_URL}.",
    "19": "Not reported: patient and public involvement or non-involvement is not stated.",
    "20a": "Results—Cohort Flow and Characteristics; Figure 1.",
    "20b": "Results—Cohort Flow and Characteristics; Table 1.",
    "20c": "Results—Cohort Flow and Characteristics; Table 1.",
    "21": "Results—Cohort Flow and Characteristics; Candidate-Model Comparison and Index-Model Selection; Model Interpretation and Locked Internal Temporal Validation.",
    "22": f"Methods—web-calculator paragraph; Code Availability; frozen model artifact and implementation code at {REPOSITORY_URL}; public calculator at {CALCULATOR_URL}.",
    "23a": "Results—Candidate-Model Comparison and Index-Model Selection; Model Interpretation and Locked Internal Temporal Validation; Figures 2 and 4; Table 2.",
    "23b": "Not applicable: cluster-level performance heterogeneity was not examined in this single-centre study.",
    "24": "Not applicable: no model updating or recalibration was performed after temporal evaluation.",
    "25": "Discussion, paragraphs 1-6. Fairness is not discussed.",
    "26": "Discussion—Limitations.",
    "27a": "Partly reported in the web implementation and Methods—web-calculator paragraph: inputs are schema-validated and invalid inputs do not produce a prediction; manuscript detail is limited.",
    "27b": "Methods—web-calculator paragraph; Discussion—clinical-implementation sentence. Clinicians enter five values; no specialist software expertise is required.",
    "27c": "Discussion—Limitations and future-research paragraph.",
}


def _insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    new_paragraph = Paragraph(element, paragraph._parent)
    new_paragraph.add_run(text)
    return new_paragraph


def _insert_before(paragraph: Paragraph, text: str, bold: bool = False) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    new_paragraph = Paragraph(element, paragraph._parent)
    run = new_paragraph.add_run(text)
    run.bold = bold
    return new_paragraph


def update_manuscript(path: Path = MANUSCRIPT_PATH) -> None:
    document = Document(path)
    for paragraph in document.paragraphs:
        if LEGACY_CALCULATOR_NOTE in paragraph.text:
            paragraph.text = paragraph.text.replace(LEGACY_CALCULATOR_NOTE, f"{CALCULATOR_NOTE}.")
        if paragraph.text.startswith("To facilitate clinical implementation") and paragraph.text.endswith(CALCULATOR_URL):
            paragraph.text = f"{paragraph.text}."
        if paragraph.text.startswith("Individual participant-level data are not publicly included"):
            paragraph.text = (
                "Participant-level data are not publicly available because of patient privacy and institutional ethics "
                "requirements."
            )
        if paragraph.text.startswith("Analysis code and the frozen five-variable model are available"):
            paragraph.text = (
                "The analysis code, frozen five-variable model artifact, and locked internal temporal-validation report "
                f"are publicly available at {REPOSITORY_URL}. The web-based research calculator is publicly accessible "
                f"at {CALCULATOR_URL}."
            )

    availability_texts = {
        "Data Availability",
        "Participant-level data are not publicly available because of patient privacy and institutional ethics requirements.",
        "Code Availability",
        (
            "The analysis code, frozen five-variable model artifact, and locked internal temporal-validation report are "
            f"publicly available at {REPOSITORY_URL}. The web-based research calculator is publicly accessible at "
            f"{CALCULATOR_URL}."
        ),
    }
    seen: set[str] = set()
    for paragraph in list(document.paragraphs):
        if paragraph.text not in availability_texts:
            continue
        if paragraph.text in seen:
            paragraph._element.getparent().remove(paragraph._element)
        else:
            seen.add(paragraph.text)
    all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if not BACKUP_PATH.exists():
        shutil.copy2(path, BACKUP_PATH)

    methods_text = (
        "To facilitate clinical implementation, the fixed Gradient Boosting specification was refitted in Group 1 using "
        "the five highest-ranked original-variable SHAP predictors (MAN-LSR, ZYG-LSR, disease duration, prior acupuncture, "
        "and age) and implemented as a web calculator that accepts these five inputs and returns the predicted probability "
        f"of one-year postoperative spasm; {CALCULATOR_NOTE}."
    )
    if "five highest-ranked original-variable SHAP predictors" not in all_text:
        target = next(p for p in document.paragraphs if p.text.startswith("A separate pooled-cohort analysis"))
        checklist_paragraph = _insert_after(target, "A completed TRIPOD+AI checklist is provided as Supplementary Table S1.")
        _insert_after(checklist_paragraph, methods_text)

    discussion_text = (
        "To support clinical implementation, we implemented the frozen five-predictor Gradient Boosting model as a web "
        "calculator for individualized risk estimation; because external validation and a clinically actionable threshold "
        "are lacking, its current use remains research-only."
    )
    if "five-predictor Gradient Boosting model as a web calculator" not in all_text:
        target = next(p for p in document.paragraphs if p.text.startswith("Decision-curve analysis demonstrated"))
        _insert_after(target, discussion_text)

    if "frozen five-variable model artifact, and locked internal temporal-validation report" not in all_text:
        target = next(p for p in document.paragraphs if p.text.strip() == "Figure Legends")
        _insert_before(target, "Data Availability", bold=True)
        _insert_before(
            target,
            "Participant-level data are not publicly available because of patient privacy and institutional ethics "
            "requirements.",
        )
        _insert_before(target, "Code Availability", bold=True)
        _insert_before(
            target,
            "The analysis code, frozen five-variable model artifact, and locked internal temporal-validation report are "
            f"publicly available at {REPOSITORY_URL}. The web-based research calculator is publicly accessible at "
            f"{CALCULATOR_URL}.",
        )

    temporary = path.with_name(f"{path.stem}.tmp{path.suffix}")
    document.save(temporary)
    try:
        os.replace(temporary, path)
    except PermissionError:
        shutil.copyfile(temporary, path)
        temporary.unlink()


def _repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def generate_checklist(path: Path = CHECKLIST_PATH) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.5)
    section.left_margin = section.right_margin = Inches(0.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TRIPOD+AI Checklist — Completed Supplementary Material")
    run.bold = True
    run.font.size = Pt(14)
    document.add_paragraph("Manuscript file: Manuscripts.docx (title and abstract are not present in the current file).")
    document.add_paragraph(
        f"Official checklist source: {SOURCE_URL}\n"
        f"TRIPOD+AI statement: {ARTICLE_URL}\n"
        f"Completed: {date.today().isoformat()}"
    )
    document.add_paragraph(
        "D = model development; E = model evaluation. Locations refer to manuscript sections because final journal pagination "
        "has not been assigned. 'Not reported', 'Not applicable', and 'Author input required' are retained explicitly rather "
        "than completed by inference."
    )

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headings = ("Section / topic", "Item", "D/E", "Checklist item", "Location in manuscript / reporting status")
    for cell, heading in zip(table.rows[0].cells, headings, strict=True):
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
    _repeat_table_header(table.rows[0])

    for item in CHECKLIST_ITEMS:
        cells = table.add_row().cells
        values = (
            f"{item.section} — {item.topic}",
            item.item,
            item.applicability,
            item.text,
            LOCATIONS[item.item],
        )
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def main() -> None:
    update_manuscript()
    print(generate_checklist())


if __name__ == "__main__":
    main()
