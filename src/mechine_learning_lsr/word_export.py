from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


RESULTS = Path("results")
SOURCE = RESULTS / "SCI_Methods_Results.md"
OUTPUT = RESULTS / "SCI_Methods_Results.docx"
SUPPLEMENTARY_LEGENDS_OUTPUT = RESULTS / "Supplementary_Figure_Legends.docx"
FIGURES = {
    number: RESULTS / "图" / filename
    for number, filename in {
        1: "Figure_1_study_flow.png",
        2: "Figure_2_model_comparison.png",
        3: "Figure_3_shap_interpretation.png",
        4: "Figure_4_temporal_validation.png",
        5: "Figure_5_exploratory_interaction.png",
    }.items()
}
TABLES = {
    number: RESULTS / "表" / f"Table_{number}_{suffix}.md"
    for number, suffix in {
        1: "baseline_characteristics",
        2: "model_comparison",
        3: "temporal_validation",
    }.items()
}
SUPPLEMENTARY_FIGURE_CAPTIONS = {
    "Figure S1. Group 1 learning curve for the Gradient Boosting model": (
        "Training AUROC and stratified five-fold cross-validated AUROC are shown as means "
        "with standard-deviation bands across five training-set sizes from 20% to 100% of "
        "the Group 1 development cohort (n=476). The plotted estimator is the fixed Gradient "
        "Boosting specification used for the index model. This learning-curve analysis is a "
        "development-cohort diagnostic and does not represent external validation. AUROC, "
        "area under the receiver operating characteristic curve."
    ),
    "Figure S2. Group 2 patient-level prediction-probability intervals": (
        "Patients in the locked Group 2 temporal-validation cohort (n=104; 9 events) are "
        "ordered by predicted probability from the frozen Gradient Boosting model. Points "
        "identify observed events and non-events, and vertical bars show 95% parameter-" 
        "uncertainty intervals obtained from 300 refits on bootstrap resamples of Group 1. "
        "The intervals describe uncertainty in the fitted model parameters, not confidence "
        "intervals for individual outcomes. No Group 2 tuning, refitting, or recalibration "
        "was performed."
    ),
    "Figure S3. Overall performance of the frozen model in Group 2": (
        "Point estimates and percentile 95% confidence intervals are shown for AUROC, AUPRC, "
        "and 1 minus the Brier score in the locked Group 2 cohort (n=104; 9 events). "
        "Intervals were estimated from 2,000 patient-level bootstrap resamples. Larger "
        "values indicate better performance for all three displayed quantities. AUROC, "
        "area under the receiver operating characteristic curve; AUPRC, area under the "
        "precision-recall curve."
    ),
    "Figure S4. Kolmogorov-Smirnov curve in Group 2": (
        "Group 2 patients are ranked from highest to lowest predicted risk. The curves show "
        "the cumulative proportions of observed events and non-events recovered along that "
        "ranking; the dashed vertical line marks the maximum separation between the curves. "
        "This is a descriptive discrimination diagnostic for the frozen model and does not "
        "establish a clinical operating threshold."
    ),
    "Figure S5. Confusion matrix for the frozen model at threshold 0.10": (
        "The matrix reports patient counts in Group 2 (n=104; 9 events) after classifying the "
        "frozen Gradient Boosting model at a predicted-probability threshold of 0.10. Rows "
        "denote observed outcomes and columns denote predicted outcomes. The threshold was "
        "used for supplementary description only and was not clinically approved or used to "
        "recalibrate or refit the model."
    ),
}
INLINE_MARKUP = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")


def _set_font(run, size: float = 12, name: str = "Times New Roman") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _add_inline(paragraph, text: str, size: float = 12) -> None:
    cursor = 0
    for match in INLINE_MARKUP.finditer(text):
        if match.start() > cursor:
            _set_font(paragraph.add_run(text[cursor : match.start()]), size)
        token = match.group()
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_font(run, size)
        elif token.startswith("`"):
            _set_font(paragraph.add_run(token[1:-1]), size, "Courier New")
        else:
            _set_font(paragraph.add_run(token[1 : token.index("]")]), size)
        cursor = match.end()
    if cursor < len(text):
        _set_font(paragraph.add_run(text[cursor:]), size)
    if "[[AUTHOR INPUT REQUIRED:" in text:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def _configure_styles(document) -> None:
    styles = document.styles
    for name, size in (("Normal", 12), ("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Title"].font.bold = True
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.bold = True
    styles["Normal"].paragraph_format.line_spacing = 2
    styles["Normal"].paragraph_format.space_after = Pt(0)


def _configure_section(section, landscape: bool = False) -> None:
    if landscape and section.page_width < section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE
    elif not landscape and section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.27)


def _add_header(document) -> None:
    paragraph = document.sections[0].header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_font(paragraph.add_run("SCI Methods and Results Draft  |  "), 9)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))
    _set_font(run, 9)


def _add_text_paragraph(document, text: str, *, compact: bool = False, quote: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1 if compact else 2
    paragraph.paragraph_format.space_after = Pt(6 if compact else 0)
    if quote:
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.right_indent = Cm(0.75)
        paragraph.paragraph_format.space_after = Pt(8)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "FFF2CC")
        paragraph._p.get_or_add_pPr().append(shading)
    _add_inline(paragraph, text, 10 if compact else 12)


def _markdown_table(path: Path) -> tuple[list[list[str]], str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    rows = []
    note = ""
    for line in lines:
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                rows.append(cells)
        elif line.startswith("**Note.**"):
            note = line
    if not rows or not note:
        raise ValueError(f"Incomplete Markdown table: {path}")
    return rows, note


def _shade_cell(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_table(document, path: Path) -> None:
    rows, note = _markdown_table(path)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    font_size = 7 if len(rows[0]) > 5 else 8.5
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.bold = row_index == 0
            _set_font(run, font_size)
            if row_index == 0:
                _shade_cell(cell, "D9EAF7")
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    _add_text_paragraph(document, note, compact=True)


def _add_picture(document, path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run().add_picture(str(path), width=Inches(8.5))


def build_word_document(source: Path = SOURCE, output: Path = OUTPUT) -> Path:
    """Convert the evidence-bounded Markdown draft into a generic SCI Word document."""
    document = Document()
    _configure_styles(document)
    _configure_section(document.sections[0])
    _add_header(document)
    document.core_properties.title = "SCI Methods and Results Draft"
    document.core_properties.subject = "Evidence-bounded Methods and Results with figures and tables"

    current_section = ""
    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("![Figure "):
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_inline(paragraph, line[2:], 16)
            continue
        if line.startswith("## "):
            current_section = line[3:]
            if current_section == "Figure Legends":
                _configure_section(document.add_section(WD_SECTION.NEW_PAGE), landscape=True)
            elif current_section == "Table Titles and Notes":
                _configure_section(document.add_section(WD_SECTION.NEW_PAGE), landscape=True)
            elif current_section == "Supplementary Material Index":
                _configure_section(document.add_section(WD_SECTION.NEW_PAGE))
            document.add_heading(current_section, level=1)
            continue
        if line.startswith("### "):
            heading = line[4:]
            figure_match = re.match(r"Figure (\d+)\.", heading)
            table_match = re.match(r"Table (\d+)\.", heading)
            if figure_match and int(figure_match.group(1)) > 1:
                document.add_page_break()
            elif table_match and int(table_match.group(1)) > 1:
                document.add_page_break()
            document.add_heading(heading, level=2)
            if figure_match:
                _add_picture(document, FIGURES[int(figure_match.group(1))])
            elif table_match:
                _add_table(document, TABLES[int(table_match.group(1))])
            continue
        if line.startswith("[Open Table "):
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, line[2:], 11)
            paragraph.paragraph_format.line_spacing = 1.15
            continue
        if line.startswith("> "):
            _add_text_paragraph(document, line[2:], quote=True)
            continue
        _add_text_paragraph(document, line, compact=current_section == "Figure Legends")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def build_supplementary_figure_legends_document(
    output: Path = SUPPLEMENTARY_LEGENDS_OUTPUT,
) -> Path:
    """Write the standalone Word document containing supplementary figure legends."""
    document = Document()
    _configure_styles(document)
    _configure_section(document.sections[0])
    _add_header(document)
    document.core_properties.title = "Supplementary Figure Legends"
    document.core_properties.subject = "Legends for supplementary figures S1-S5"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(title.add_run("Supplementary Figure Legends"), 16)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(subtitle.add_run("Figures S1-S5"), 11)

    for heading, caption in SUPPLEMENTARY_FIGURE_CAPTIONS.items():
        document.add_heading(heading, level=1)
        _add_text_paragraph(document, caption)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main() -> None:
    output = build_word_document()
    print(f"Generated {output} with five figures and three editable tables.")


if __name__ == "__main__":
    main()
